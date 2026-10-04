import re
import sys
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def sanitize_xml(s):
    if not s:
        return ''
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]', '', str(s))

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', 'auto'))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def add_formatted_runs(paragraph, text):
    text = sanitize_xml(text)
    # Tokenize: `code`, **bold**, *italic*
    pattern = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)')
    tokens = pattern.split(text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            r = paragraph.add_run(sanitize_xml(token[2:-2]))
            r.bold = True
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            r = paragraph.add_run(sanitize_xml(token[1:-1]))
            r.italic = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            r = paragraph.add_run(sanitize_xml(token[1:-1]))
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(160, 40, 40)
        else:
            paragraph.add_run(sanitize_xml(token))

def convert_markdown_to_docx(md_path, docx_path):
    print(f"Convirtiendo {md_path.name} -> {docx_path.name}...")
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    lines = md_path.read_text(encoding='utf-8').splitlines()
    i = 0
    n = len(lines)
    
    in_table = False
    table_lines = []
    
    while i < n:
        line = lines[i]
        stripped = line.strip()
        
        # Table detection
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines.append(stripped)
            i += 1
            continue
        elif table_lines:
            process_table(doc, table_lines)
            table_lines = []
        
        if not stripped:
            i += 1
            continue
            
        # Horizontal Rule
        if stripped in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run('—' * 55)
            r.font.color.rgb = RGBColor(203, 213, 225)
            r.bold = True
            i += 1
            continue
            
        # Headings
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(sanitize_xml(stripped[2:]))
            r.font.name = 'Calibri'
            r.font.size = Pt(18)
            r.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue
            
        if stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(sanitize_xml(stripped[3:]))
            r.font.name = 'Calibri'
            r.font.size = Pt(14)
            r.bold = True
            r.font.color.rgb = RGBColor(30, 58, 138)
            i += 1
            continue

        if stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(sanitize_xml(stripped[4:]))
            r.font.name = 'Calibri'
            r.font.size = Pt(12)
            r.bold = True
            r.font.color.rgb = RGBColor(3, 105, 161)
            i += 1
            continue

        if stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(sanitize_xml(stripped[5:]))
            r.font.name = 'Calibri'
            r.font.size = Pt(11)
            r.bold = True
            r.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue

        # Blockquote (> ...)
        if stripped.startswith('>'):
            bq_lines = []
            while i < n and lines[i].strip().startswith('>'):
                bq_lines.append(lines[i].strip()[1:].strip())
                i += 1
            process_blockquote(doc, bq_lines)
            continue

        # Bullet list
        if stripped.startswith(('* ', '- ', '• ')):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            content = stripped[2:].strip()
            add_formatted_runs(p, content)
            i += 1
            continue

        # Numbered list
        m_num = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m_num:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            content = m_num.group(2).strip()
            add_formatted_runs(p, content)
            i += 1
            continue

        # Regular Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_runs(p, stripped)
        i += 1

    if table_lines:
        process_table(doc, table_lines)

    doc.save(docx_path)
    print(f"  -> Guardado exitosamente: {docx_path}")

def process_blockquote(doc, bq_lines):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=100, bottom=100, left=180, right=140)
    set_cell_border(cell, left=dict(sz=24, val='single', color='0284C7'),
                          top=dict(val='none'), right=dict(val='none'), bottom=dict(val='none'))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    for idx, line in enumerate(bq_lines):
        if idx > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
        add_formatted_runs(p, line)
    
    sp_p = doc.add_paragraph()
    sp_p.paragraph_format.space_before = Pt(0)
    sp_p.paragraph_format.space_after = Pt(4)

def process_table(doc, lines):
    rows = []
    for l in lines:
        cells = [c.strip() for c in l.split('|')]
        if cells and cells[0] == '':
            cells = cells[1:]
        if cells and cells[-1] == '':
            cells = cells[:-1]
        if all(re.match(r'^:?-+:?$', c) for c in cells):
            continue
        rows.append(cells)
    
    if not rows:
        return
        
    num_cols = max(len(r) for r in rows)
    for r in rows:
        while len(r) < num_cols:
            r.append('')
            
    tbl = doc.add_table(rows=len(rows), cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    
    for r_idx, row_data in enumerate(rows):
        row = tbl.rows[r_idx]
        is_header = (r_idx == 0)
        
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            
        for c_idx, cell_value in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            
            if is_header:
                set_cell_background(cell, "1E3A8A")
                set_cell_border(cell, top=dict(sz=6, color='1E3A8A'),
                                      bottom=dict(sz=12, color='1E3A8A'),
                                      left=dict(sz=4, color='CBD5E1'),
                                      right=dict(sz=4, color='CBD5E1'))
                add_formatted_runs(p, cell_value)
                for run in p.runs:
                    run.font.bold = True
                    run.font.size = Pt(9.0)
                    run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                bg = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
                set_cell_background(cell, bg)
                set_cell_border(cell, top=dict(sz=4, color='E2E8F0'),
                                      bottom=dict(sz=4, color='E2E8F0'),
                                      left=dict(sz=4, color='E2E8F0'),
                                      right=dict(sz=4, color='E2E8F0'))
                add_formatted_runs(p, cell_value)
                for run in p.runs:
                    run.font.size = Pt(8.5)
                    run.font.name = 'Calibri'
                    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(2)
    p_after.paragraph_format.space_after = Pt(6)

if __name__ == '__main__':
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2])
    convert_markdown_to_docx(src, dst)
